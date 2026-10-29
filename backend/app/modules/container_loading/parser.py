"""Parser for the 'DSR Report Template - Container Loading Monitoring' Word document.

Design rules (graceful termination):
  * parse() never raises for bad input - every problem becomes an Issue.
  * ERROR issues mean the file cannot be trusted -> the upload is aborted and nothing is saved.
  * WARNING issues are saved with the data and shown in the upload log.
  * All problems are collected in one pass so the user can fix the document once.
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..base import Issue, ParseResult
from ..docx_utils import (clean, is_blank, iter_blocks, kv_get, kv_of, norm_key, parse_date, parse_dims,  # noqa: F401
                          parse_float, parse_int, row_cells, strip_colon)

MAX_CONTAINERS = 5000
MAX_PHOTO_BYTES = 8 * 1024 * 1024

CHECKLIST_FIELDS = {  # output key -> label prefix in "Summary of Inspection"
    "weather": "weather conditions",
    "storage": "storage conditions",
    "packing": "packing quality",
    "loading": "loading quality",
    "packingListValidation": "validation of packing list",
    "outgoingQc": "outgoing quality control",
    "shockLabel": "shock label",
}

CONTAINER_COLUMNS = {  # field -> predicate on normalised header text
    "srNo": lambda h: h == "no",
    "sealNo": lambda h: "seal" in h,
    "containerNo": lambda h: h in ("containers no", "container no", "container number"),
    "palletHuNo": lambda h: "pallet" in h,
    "modelCode": lambda h: "product type" in h,
    "nominalPowerWp": lambda h: "nominal power" in h,
    "currentClass": lambda h: "current class" in h,
    "moduleQty": lambda h: "module quantity" in h,
    "containerSize": lambda h: "container size" in h,
    "loadingDate": lambda h: "date of loading" in h,
}
REQUIRED_COLUMNS = ("containerNo", "modelCode", "nominalPowerWp", "moduleQty", "loadingDate")


def status_of(s: str, item: str) -> str:
    """Normalise a checklist cell to OK / NOK / NA, or '?' if unrecognised."""
    v = norm_key(s)
    if v in ("ok", "okay", "pass", "passed", "yes", "available"):
        return "OK"
    if v in ("nok", "not ok", "fail", "failed"):
        return "NOK"
    if v in ("", "na", "n a", "nil", "not applicable") or is_blank(s):
        return "NA"
    if v == "no" and item == "shockLabel":  # "Shock Label Validations (if applicable): No" => not applicable
        return "NA"
    return "?"


# ---------- container table ----------
def _header_map(table: Table) -> tuple[dict[str, int], int]:
    """Return ({field: column index}, index of first data row). Header text is joined across header rows."""
    rows = table.rows
    first_data = None
    for i, r in enumerate(rows):
        cells = row_cells(r)
        if cells and re.fullmatch(r"\d+", cells[0]) and i > 0:
            first_data = i
            break
    if first_data is None:
        return {}, len(rows)
    ncol = len(rows[0].cells)
    headers = []
    for j in range(ncol):
        parts: list[str] = []
        for i in range(first_data):
            t = clean(rows[i].cells[j].text)
            if t and (not parts or parts[-1] != t):
                parts.append(t)
        headers.append(norm_key(" ".join(parts)))
    mapping: dict[str, int] = {}
    for field, pred in CONTAINER_COLUMNS.items():
        for j, h in enumerate(headers):
            if j not in mapping.values() and pred(h):
                mapping[field] = j
                break
    return mapping, first_data


def _find_tables(doc):
    containers = summary = None
    for t in doc.tables:
        head = " ".join(norm_key(" ".join(row_cells(r))) for r in t.rows[:3])
        if "containers no" in head and "nominal power" in head:
            containers = containers or t
        elif "psi lot" in head and "date of loading" in head and "no of containers" in head:
            summary = summary or t
    return containers, summary


def _parse_summary(table: Table, issues: list[Issue]) -> list[dict[str, Any]]:
    rows = table.rows
    start = None
    for i, r in enumerate(rows):
        c = row_cells(r)
        if c and parse_date(c[0]):
            start = i
            break
    if start is None:
        issues.append(Issue("WARNING", "SUMMARY_EMPTY", "Inspection-summary table has no data rows; PSI lots cannot be assigned."))
        return []
    hdr = [norm_key(" ".join(dict.fromkeys(clean(rows[i].cells[j].text) for i in range(start)))) for j in range(len(rows[0].cells))]
    idx = {"date": 0, "type": 1, "lot": 2, "containers": 3}
    for j, h in enumerate(hdr):
        if "date of loading" in h:
            idx["date"] = j
        elif "product type" in h:
            idx["type"] = j
        elif "psi lot" in h:
            idx["lot"] = j
    # first "total" column under "No. of Containers"
    for j, h in enumerate(hdr):
        if "no of containers" in h and h.endswith("total"):
            idx["containers"] = j
            break
    for j, h in enumerate(hdr):
        if "no of modules" in h and h.endswith("total"):
            idx["modules"] = j
            break
    out = []
    for r in rows[start:]:
        c = row_cells(r)
        d = parse_date(c[idx["date"]]) if len(c) > idx["date"] else None
        if not d:
            continue
        out.append({"loadingDate": d, "modelCode": c[idx["type"]], "lotNo": c[idx["lot"]],
                    "containers": parse_int(c[idx["containers"]]) if len(c) > idx["containers"] else None,
                    "modules": parse_int(c[idx["modules"]]) if "modules" in idx and len(c) > idx["modules"] else None})
    return out


# ---------- photos ----------
PHOTO_HINTS = ("container seal", "sample pallet picture", "empty container", "loading process", "closed container",
               "container after loading", "transporter packing list", "vehicle number", "fully loaded container",
               "invoice and out time", "container closing")


def _photos_from_table(doc, table: Table, issues: list[Issue], sr: int) -> list[dict]:
    photos = []
    rows = table.rows
    for i in range(len(rows) - 1):
        for j, cell in enumerate(rows[i].cells):
            label = clean(cell.text)
            if not label:
                continue
            below = rows[i + 1].cells[j]
            for rid in below._tc.xpath(".//a:blip/@r:embed"):
                try:
                    part = doc.part.related_parts[rid]
                    blob = part.blob
                except Exception:  # noqa: BLE001
                    issues.append(Issue("WARNING", "PHOTO_UNREADABLE", f"Container {sr}: photo '{label}' could not be read.", {"container": sr}))
                    continue
                if len(blob) > MAX_PHOTO_BYTES:
                    issues.append(Issue("WARNING", "PHOTO_TOO_LARGE", f"Container {sr}: photo '{label}' skipped (over 8 MB).", {"container": sr}))
                    continue
                ext = part.partname.ext.lstrip(".").lower() or "png"
                photos.append({"label": label, "ext": ext, "contentType": part.content_type, "data": blob})
    return photos


# ---------- main entry ----------
def parse_container_loading(data: bytes) -> ParseResult:
    issues: list[Issue] = []

    def err(code, msg, **ctx):
        issues.append(Issue("ERROR", code, msg, ctx))

    def warn(code, msg, **ctx):
        issues.append(Issue("WARNING", code, msg, ctx))

    # 1. container-level checks
    src = io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else str(data)  # bytes or a path on disk
    if not zipfile.is_zipfile(src):
        err("NOT_DOCX", "The file is not a valid .docx document (legacy .doc, corrupted, or renamed file).")
        return ParseResult(None, issues)
    if isinstance(src, io.BytesIO):
        src.seek(0)
    try:
        doc = Document(src)
    except Exception as e:  # noqa: BLE001
        err("DOCX_UNREADABLE", f"The Word document could not be opened: {type(e).__name__}.")
        return ParseResult(None, issues)

    # 2. report header
    report: dict[str, Any] = {"client": {}, "signatories": []}
    kv_all: dict[str, str] = {}
    for t in doc.tables:
        kv = kv_of(t)
        if "report no" in kv:
            kv_all = kv
        if "spv name" in kv or "project name" in kv:  # newer template: Client / Project Name / SPV Name table
            report["client"] = {"name": kv.get("client", ""), "address": ""}
            report["spvName"], report["projectName"] = kv.get("spv name", ""), kv.get("project name", "")
        first = clean(t.rows[0].cells[0].text) if t.rows else ""
        if norm_key(first) == "client details" and len(t.rows) > 1:
            lines = [clean(x) for x in t.rows[1].cells[0].text.splitlines() if clean(x)]
            report["client"] = {"name": lines[0] if lines else "", "address": ", ".join(lines[1:])}
        elif norm_key(first) == "inspection conclusion" and len(t.rows) > 1:
            report["conclusion_raw"] = clean(t.rows[1].cells[0].text)
        elif norm_key(first) in ("prepared by", "reviewed by", "approved by"):
            for r in t.rows:
                c = row_cells(r)
                report["signatories"].append({"role": c[0], "designation": c[1] if len(c) > 1 else "",
                                              "name": c[2] if len(c) > 2 else ""})
    if not kv_all:
        err("REPORT_INFO_MISSING", "The 'Report No. / Report Date / Manufacturer / Location / Workshop' table was not found.")
        return ParseResult(None, issues)

    report_no = kv_all.get("report no", "")
    if is_blank(report_no):
        err("REPORT_NO_MISSING", "Report No. is empty or still a placeholder ('xxxxx'). Fill it in before uploading.")
    report["reportNo"] = report_no
    rd = parse_date(kv_all.get("report date", ""))
    if not rd:
        err("REPORT_DATE_INVALID", f"Report Date '{kv_all.get('report date', '')}' is not a valid date (expected e.g. 16-09-2026).")
    report["reportDate"] = rd
    for k, label in (("manufacturer", "Manufacturer"), ("location", "Location"), ("workshop", "Workshop")):
        v = kv_all.get(k, "")
        if is_blank(v):
            warn("HEADER_FIELD_BLANK", f"{label} is blank or a placeholder.")
            v = "Unknown" if k == "manufacturer" else ""
        report[k] = v
    if not report["client"].get("name"):
        warn("CLIENT_MISSING", "Client details not found.")
    conc = norm_key(report.pop("conclusion_raw", ""))
    report["conclusion"] = "PASS" if conc == "pass" else "FAIL" if conc == "fail" else None
    if report["conclusion"] is None:
        warn("CONCLUSION_UNSET", "Inspection conclusion is not exactly 'Pass' or 'Fail'; it will be derived from container results.")

    # 3. tables
    ctable, stable = _find_tables(doc)
    if ctable is None:
        err("CONTAINER_TABLE_MISSING", "The 'Containers' details' table (Containers No., Nominal Power, ...) was not found.")
        return ParseResult(None, issues)
    mapping, first_data = _header_map(ctable)
    missing = [f for f in REQUIRED_COLUMNS if f not in mapping]
    if missing:
        err("CONTAINER_COLUMNS_MISSING", f"Containers table is missing required column(s): {', '.join(missing)}.", columns=missing)
        return ParseResult(None, issues)

    summary_rows = _parse_summary(stable, issues) if stable is not None else []
    if stable is None:
        warn("SUMMARY_TABLE_MISSING", "General-information summary table not found; PSI lots cannot be assigned.")
    lot_by_key = {(s["loadingDate"].date(), norm_key(s["modelCode"])): s["lotNo"] for s in summary_rows}

    # 4. containers
    containers: list[dict[str, Any]] = []
    lot_missing: dict[tuple, list[str]] = {}
    seen_nos: set[str] = set()
    for ri in range(first_data, len(ctable.rows)):
        cells = row_cells(ctable.rows[ri])
        if not any(cells):
            continue
        if len(containers) >= MAX_CONTAINERS:
            err("TOO_MANY_CONTAINERS", f"More than {MAX_CONTAINERS} containers in one report.")
            break
        g = {f: (cells[j] if j < len(cells) else "") for f, j in mapping.items()}
        label = f"row {ri + 1} (container '{g.get('containerNo') or '?'}')"
        sr = parse_int(g.get("srNo", "")) or len(containers) + 1
        ok = True
        cno = g["containerNo"].upper().replace(" ", "")
        if is_blank(cno):
            err("CONTAINER_NO_MISSING", f"Containers table {label}: container number is empty.", row=ri + 1); ok = False
        elif cno in seen_nos:
            err("CONTAINER_NO_DUPLICATE", f"Container number {cno} appears more than once in this report.", container=cno); ok = False
        seen_nos.add(cno)
        qty = parse_int(g["moduleQty"])
        if qty is None or qty <= 0:
            err("MODULE_QTY_INVALID", f"Containers table {label}: module quantity '{g['moduleQty']}' is not a positive number.", row=ri + 1); ok = False
        power = parse_float(g["nominalPowerWp"])
        if not power or power <= 0:
            m = re.search(r"(?:^|[-_])(\d{3})[A-Z]", g["modelCode"].upper())
            power = float(m.group(1)) if m else None
            if power:
                warn("POWER_FROM_MODEL", f"{cno}: nominal power missing; taken from model code ({power:g} Wp).", container=cno)
            else:
                err("POWER_INVALID", f"Containers table {label}: nominal power '{g['nominalPowerWp']}' is invalid.", row=ri + 1); ok = False
        ld = parse_date(g["loadingDate"])
        if not ld:
            err("LOADING_DATE_INVALID", f"Containers table {label}: loading date '{g['loadingDate']}' is not a valid date.", row=ri + 1); ok = False
        if is_blank(g["modelCode"]):
            err("PRODUCT_TYPE_MISSING", f"Containers table {label}: product type is empty.", row=ri + 1); ok = False
        if not ok:
            continue
        model = g["modelCode"].upper().strip()
        lot = lot_by_key.get((ld.date(), norm_key(model)))
        if lot is None:
            lot_missing.setdefault((ld.date(), model), []).append(cno)
            lot = "N/A"
        containers.append({
            "srNo": sr, "containerNo": cno, "sealNo": g.get("sealNo", ""), "palletHuNo": g.get("palletHuNo", ""),
            "modelCode": model, "nominalPowerWp": power, "currentClass": g.get("currentClass", ""),
            "moduleQty": qty, "containerSize": g.get("containerSize", ""), "loadingDate": ld,
            "capacityKwp": round(qty * power / 1000, 3), "psiLotNo": lot,
            "psiLots": [x for x in re.split(r"[,&/;\s]+", lot) if x] or ["N/A"],  # "1, 2" => ["1", "2"]
            "shipment": {}, "inspection": {}, "photos": [],
        })
    for (d, model), nos in lot_missing.items():
        warn("LOT_NOT_FOUND", f"No PSI lot in the summary table for {d:%d-%m-%Y} / {model} ({len(nos)} container(s)); lot set to 'N/A'.", containers=nos)
    if not containers and not errors_exist(issues):
        err("NO_CONTAINERS", "The containers table has no data rows.")
    if any(i.level == "ERROR" for i in issues):
        return ParseResult(None, issues)

    # 5. per-container sections (shipment / inspection / photos)
    by_sr = {c["srNo"]: c for c in containers}
    current: int | None = None
    sections_seen: set[int] = set()
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            m = re.match(r"\s*(?:shipment details\s*[-–—:]\s*)?container number\s*(\d+)\s*$", block.text, re.I)
            if m:
                current = int(m.group(1))
                sections_seen.add(current)
                if current not in by_sr:
                    warn("SECTION_NO_CONTAINER", f"'Container Number {current}' section has no matching row in the containers table; ignored.")
            continue
        text = norm_key(" ".join(" ".join(row_cells(r)) for r in block.rows))
        kv = kv_of(block)
        target = by_sr.get(current) if current is not None else None
        if "warehousing site" in kv:
            if target is None:
                continue
            target["shipment"] = _shipment(kv, target, warn, parse_date)
        elif "inspection result" in kv:
            if target is None:
                continue
            target["inspection"] = _inspection(kv, target, warn)
        elif any(h in text for h in PHOTO_HINTS):
            if target is not None:
                target["photos"] += _photos_from_table(doc, block, issues, current)
    no_section = [c["containerNo"] for c in containers if c["srNo"] not in sections_seen]
    if no_section:
        warn("SECTION_MISSING", f"{len(no_section)} container(s) have no 'Shipment Details' section (inspection data unavailable): "
             f"{', '.join(no_section[:8])}{'...' if len(no_section) > 8 else ''}.", containers=no_section)
    for c in containers:
        if c["srNo"] in sections_seen and not c["inspection"]:
            warn("INSPECTION_MISSING", f"Container {c['containerNo']}: 'Summary of Inspection' table not found.", container=c["containerNo"])
        if not c["inspection"]:
            c["inspection"] = {"result": None, "checklist": {}}

    # 6. cross-checks against the document's own summary table
    for s in summary_rows:
        match = [c for c in containers if c["loadingDate"].date() == s["loadingDate"].date() and norm_key(c["modelCode"]) == norm_key(s["modelCode"])]
        n, mods = len(match), sum(c["moduleQty"] for c in match)
        if s["containers"] is not None and n != s["containers"]:
            warn("SUMMARY_MISMATCH", f"Summary table says {s['containers']} container(s) for {s['loadingDate']:%d-%m-%Y} / {s['modelCode']} "
                 f"but the containers table lists {n}. Dashboard uses the containers table.")
        if s.get("modules") is not None and mods != s["modules"]:
            warn("SUMMARY_MODULES_MISMATCH", f"Summary table says {s['modules']:,} modules for {s['loadingDate']:%d-%m-%Y} / {s['modelCode']} "
                 f"but the containers table adds up to {mods:,}. Dashboard uses the containers table.")
    results = [c["inspection"].get("result") for c in containers]
    if report["conclusion"] is None:
        report["conclusion"] = "FAIL" if "FAIL" in results else "PASS" if "PASS" in results else None
    elif report["conclusion"] == "PASS" and "FAIL" in results:
        warn("CONCLUSION_CONFLICT", "Report conclusion is PASS but at least one container inspection FAILED.")
    report["lots"] = summary_rows
    issues.insert(0, Issue("INFO", "PARSED", f"Parsed report {report_no}: {len(containers)} container(s), "
                           f"{sum(c['moduleQty'] for c in containers):,} modules."))
    return ParseResult({"report": report, "containers": containers}, issues)


def errors_exist(issues: list[Issue]) -> bool:
    return any(i.level == "ERROR" for i in issues)


def _shipment(kv, c, warn, parse_date):
    seal = kv_get(kv, "seal number") or ""
    if re.fullmatch(r"(n/?a|nil|-+)?", seal.strip(), re.I) or is_blank(seal):
        seal = ""  # "N/A" is not a seal number
    if seal and not c["sealNo"]:
        c["sealNo"] = seal  # newer template has no Seal column in the containers table
    elif seal and clean(seal).upper() != c["sealNo"].upper():
        warn("SEAL_MISMATCH", f"Container {c['containerNo']}: seal in shipment details ('{seal}') differs from containers table ('{c['sealNo']}'). Containers table used.",
             container=c["containerNo"])
    dep = kv_get(kv, "time of truck departure") or ""
    return {
        "warehousingSite": kv_get(kv, "warehousing site"), "containerType": kv_get(kv, "container type"),
        "loadingStatus": kv_get(kv, "container loading status"), "modulePosition": kv_get(kv, "module position"),
        "truckPlateNo": kv_get(kv, "truck plate"), "truckDeparture": parse_date(dep) or (dep or None),
        "noOfProducts": parse_int(kv_get(kv, "no of products") or ""), "fullPallets": parse_int(kv_get(kv, "no of full pallets") or ""),
        "partialPallets": parse_int(kv_get(kv, "number of partially") or ""),
    }


def _inspection(kv, c, warn):
    res_raw = norm_key(kv_get(kv, "inspection result") or "")
    result = "PASS" if res_raw in ("pass", "passed") else "FAIL" if res_raw in ("fail", "failed") else None
    if result is None:
        warn("RESULT_UNKNOWN", f"Container {c['containerNo']}: inspection result '{res_raw}' is not PASS/FAIL.", container=c["containerNo"])
    checklist = {}
    for key, prefix in CHECKLIST_FIELDS.items():
        raw = kv_get(kv, prefix) or ""
        st = status_of(raw, key)
        if st == "?":
            warn("CHECKLIST_VALUE", f"Container {c['containerNo']}: '{prefix}' value '{raw}' not recognised; treated as NOK.", container=c["containerNo"])
            st = "NOK"
        checklist[key] = st
    pallet = parse_dims(kv_get(kv, "pallet dimension") or "")
    mod = parse_dims(kv_get(kv, "module dimension") or "")
    nc = kv_get(kv, "no of non conformity") or ""
    return {
        "result": result, "inspectionQty": parse_int(kv_get(kv, "inspection quantity") or ""),
        "nonConformityCount": 0 if norm_key(nc) in ("nil", "none", "", "0") else (parse_int(nc) or 1),
        "observation": kv_get(kv, "observation"),
        "pallet": {"type": kv_get(kv, "pallet type"), "dimMm": pallet, "volM3": parse_float(kv_get(kv, "pallet volume") or "")},
        "instructionManual": kv_get(kv, "instruction manual"), "moduleDimMm": mod, "checklist": checklist,
    }
