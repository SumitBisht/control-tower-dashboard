"""End-to-end API test against a real MongoDB (skipped if none is reachable). Uses a throw-away database."""
import asyncio
import os
import uuid

import pymongo
import pytest
from httpx import ASGITransport, AsyncClient
from pathlib import Path

URI = os.environ.get("MONGODB_URI", "mongodb://localhost:27017")
DB = f"ct_test_{uuid.uuid4().hex[:8]}"
os.environ.update({"MONGODB_DB": DB, "MONGODB_URI": URI, "STORAGE_LOCAL_DIR": f"./.test_storage/{DB}",
                   "ADMIN_EMAIL": "admin@test.io", "ADMIN_PASSWORD": "AdminPass123"})

FIX = Path(__file__).parent / "fixtures"


def _mongo_up() -> bool:
    try:
        pymongo.MongoClient(URI, serverSelectionTimeoutMS=1500).admin.command("ping")
        return True
    except Exception:  # noqa: BLE001
        return False


pytestmark = pytest.mark.skipif(not _mongo_up(), reason="MongoDB not reachable")


@pytest.fixture
async def client():
    from app.core.db import close_db, get_db
    from app.main import app, bootstrap

    await bootstrap(get_db())
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t/api") as c:
        yield c
    await get_db().client.drop_database(DB)
    await close_db()


async def login(c, email, pw):
    r = await c.post("/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['accessToken']}"}


async def upload_and_wait(c, headers, name, data, **form):
    r = await c.post("/uploads", headers=headers, files={"file": (name, data)}, data=form)
    assert r.status_code == 202, r.text
    uid = r.json()["id"]
    for _ in range(300):
        g = (await c.get(f"/uploads/{uid}", headers=headers)).json()
        if g["status"] != "PROCESSING":
            return g
        await asyncio.sleep(0.1)
    raise AssertionError("upload did not finish")


async def test_full_flow(client):
    c = client
    admin = await login(c, "admin@test.io", "AdminPass123")

    # self-registration -> pending -> cannot log in -> admin approves
    assert (await c.post("/auth/register", json={"name": "V", "email": "v@t.io", "password": "password1"})).status_code == 201
    assert (await c.post("/auth/login", json={"email": "v@t.io", "password": "password1"})).status_code == 403
    uid = next(u["id"] for u in (await c.get("/admin/users", headers=admin)).json() if u["email"] == "v@t.io")
    assert (await c.post(f"/admin/users/{uid}/approve", headers=admin)).status_code == 200
    viewer = await login(c, "v@t.io", "password1")

    # roles
    assert (await c.get("/admin/users", headers=viewer)).status_code == 403
    r = await c.post("/uploads", headers=viewer, files={"file": ("a.docx", b"x")})
    assert r.status_code == 403
    assert (await c.post("/admin/users", headers=admin, json={"name": "U", "email": "u@t.io", "password": "password1", "role": "uploader"})).status_code == 201
    uploader = await login(c, "u@t.io", "password1")

    # last admin protection
    me = (await c.get("/auth/me", headers=admin)).json()
    assert (await c.patch(f"/admin/users/{me['id']}", headers=admin, json={"role": "viewer"})).status_code == 400

    # failed uploads terminate gracefully and save nothing
    for name, data in [("x.pdf", b"%PDF"), ("bad.docx", b"not a zip"), ("tpl.docx", (FIX / "template.docx").read_bytes())]:
        r = await upload_and_wait(c, uploader, name, data)
        assert r["status"] == "FAILED", name
        assert any(line["level"] == "ERROR" for line in r["logs"])
    assert (await c.get("/container_loading/dashboard", headers=viewer)).json()["widgets"]["kpi_containers"]["value"] == 0

    # good upload
    good = (FIX / "sample_valid.docx").read_bytes()
    r = await upload_and_wait(c, uploader, "ok.docx", good, project_name="Proj A")
    assert r["status"] == "SUCCESS_WITH_WARNINGS", r["logs"]
    assert (await upload_and_wait(c, uploader, "ok.docx", good))["status"] == "FAILED"  # duplicate
    assert (await upload_and_wait(c, uploader, "ok.docx", good, replace="true"))["status"] != "FAILED"

    w = (await c.get("/container_loading/dashboard", headers=viewer)).json()["widgets"]
    assert w["kpi_containers"]["value"] == 12 and w["kpi_modules"]["value"] == 8640
    assert w["kpi_capacity"]["value"] == 4.752
    assert w["product_summary"]["rows"][0]["containers"] == 12
    assert w["loading_trend"]["cumulative"] and w["loading_trend"]["data"][-1]["550 Wp"] == 8640
    assert {r["item"]: r["status"] for r in w["checklist_summary"]["rows"]}["Shock Label Validations"] == "NA"
    assert (await c.get("/container_loading/dashboard?project=Nope", headers=viewer)).json()["widgets"]["kpi_containers"]["value"] == 0

    # widget + column hiding is enforced server-side
    assert (await c.put("/admin/widgets/container_loading/kpi_modules", headers=admin, json={"visible": False})).status_code == 200
    assert (await c.put("/admin/widgets/container_loading/product_summary", headers=admin, json={"columns": {"modules": False}})).status_code == 200
    w = (await c.get("/container_loading/dashboard", headers=viewer)).json()["widgets"]
    assert "kpi_modules" not in w
    assert "modules" not in w["product_summary"]["rows"][0]
    assert (await c.put("/admin/widgets/container_loading/product_summary", headers=admin, json={"columns": {"bogus": False}})).status_code == 422


async def test_premier_report_dashboard_and_reports_tab(client):
    c = client
    admin = await login(c, "admin@test.io", "AdminPass123")
    raw = (FIX / "premier_report_slim.docx").read_bytes()
    up = await upload_and_wait(c, admin, "premier.docx", raw)
    assert up["status"] in ("SUCCESS", "SUCCESS_WITH_WARNINGS"), up["logs"]
    assert up["summary"]["containers"] == 25 and up["summary"]["modules"] == 15066 and up["summary"]["photos"] == 196
    assert up["summary"]["spvName"].startswith("Solapur Ganeko One") and up["summary"]["projectName"].startswith("Ganeko Two")

    w = (await c.get("/container_loading/dashboard", headers=admin)).json()["widgets"]
    assert w["kpi_containers"]["value"] == 25 and w["kpi_modules"]["value"] == 15066 and w["kpi_capacity"]["value"] == 9.266
    assert w["kpi_result"]["value"] == 100.0
    assert w["loading_trend"]["data"][-1]["615 Wp"] == 15066 and len(w["loading_trend"]["data"]) == 3
    f = (await c.get("/container_loading/filters", headers=admin)).json()
    assert f["psi_lot"] == ["1", "2"] and f["model"] == ["PE-132-615THGB-G12R"]
    # PSI lot "1, 2" containers belong to both lots
    lot1 = (await c.get("/container_loading/dashboard?psi_lot=1", headers=admin)).json()["widgets"]["kpi_containers"]["value"]
    lot2 = (await c.get("/container_loading/dashboard?psi_lot=2", headers=admin)).json()["widgets"]["kpi_containers"]["value"]
    assert lot1 + lot2 > 25 and lot1 < 25 and lot2 < 25

    # reports tab: filters
    rep = (await c.get("/container_loading/reports", headers=admin)).json()
    assert rep["total"] == 1 and rep["items"][0]["containers"] == 25 and rep["items"][0]["hasSource"]
    base = "/container_loading/reports?"
    assert (await c.get(base + "date_from=2026-10-09&date_to=2026-10-09", headers=admin)).json()["items"][0]["containers"] == 8
    assert (await c.get(base + "q=MH14KQ5656", headers=admin)).json()["items"][0]["containers"] == 1
    assert (await c.get(base + "q=nomatch", headers=admin)).json()["total"] == 0
    assert (await c.get(base + "model=PE-132-615THGB-G12R&date_from=2026-11-01", headers=admin)).json()["total"] == 0
    assert (await c.get(base + "date_from=2026-10-10&date_to=2026-10-01", headers=admin)).status_code == 422
    # downloads
    d = await c.get(f"/container_loading/reports/{rep['items'][0]['id']}/download", headers=admin)
    assert d.status_code == 200 and d.content == raw
    x = await c.get("/container_loading/reports/export?format=xlsx&date_from=2026-10-10", headers=admin)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook
    import io
    rows = list(load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True))
    assert rows[0][0] == "Report No" and len(rows) == 1 + 4
    csv_ = await c.get("/container_loading/reports/export?format=csv&q=MH46BU7742", headers=admin)
    assert "MH46BU7742" in csv_.text and len(csv_.text.strip().splitlines()) == 2
    assert (await c.get("/container_loading/reports/export?q=nomatch", headers=admin)).status_code == 404
