"""Pre-Shipment Inspection: parser tests (always run) + end-to-end API test (needs MongoDB)."""
import io
import os
from pathlib import Path

import pytest
from docx import Document

from app.modules.pre_shipment.parser import parse_pre_shipment

FIX = Path(__file__).parent / "fixtures"
L1, L2 = "psi_lot-1.docx", "psi_lot-2.docx"


def parse(name):
    return parse_pre_shipment((FIX / name).read_bytes())


def codes(r, level):
    return {i.code for i in r.issues if i.level == level}


def test_lot_1():
    r = parse(L1)
    assert r.ok, [i.message for i in r.errors]
    d = r.data
    assert d["report"]["reportNo"] == "DSR-PSI-RES-2026-10-04" and d["report"]["reportDate"].day == 4
    assert d["report"]["spvName"].startswith("Solapur Ganeko One") and d["report"]["projectName"].startswith("Ganeko Two")
    assert d["lot"] == {"lotNo": "01", "factory": "P7 Hyderabad", "sizePcs": 8246, "sizeMwp": 5.071}
    assert d["wattClasses"][0]["wattClass"] == 615 and d["wattClasses"][0]["modules"] == 8246
    assert d["conclusion"] == "PASS" and len(d["results"]) == 12 and all(x["decision"] == "PASS" for x in d["results"])
    assert d["ncs"] == [{"test": "Visual Check", "category": "Minor", "description": "Frame Scratch", "risk": "Minor", "count": 1}]
    mp = d["details"]["maxPower"]
    assert len(mp["rows"]) == 80 and set(mp["stats"]) == {"minimum", "maximum", "average"} and mp["summary"]["outOfTolerance"] == 0
    assert mp["rows"][0]["moduleNo"] == "NSME26090116872" and mp["rows"][0]["diffPct"] == -0.31
    assert len(d["details"]["testConditions"]) == 10 and d["details"]["testConditions"][0]["nextCalibration"].year == 2027
    assert len(d["details"]["calibration"]) == 4 and d["details"]["calibration"][0]["test_time"].hour == 16
    det = d["details"]
    assert len(det["insulation"]) == 5 and len(det["wetLeakage"]) == 5 and len(det["biFaciality"]) == 2 and len(det["lowIrradiance"]) == 2
    assert len(det["dynamicLoad"]) == 1 and len(det["staticLoad"]) == 1 and len(det["robustness"]) == 1
    assert det["ncSummary"]["Visual Check"][0]["observedMinor"] == 1
    assert not codes(r, "WARNING")  # a clean report


def test_lot_2_flags_sampling_shortfall_and_keeps_stat_order():
    r = parse(L2)
    assert r.ok
    assert r.data["lot"]["lotNo"] == "02" and r.data["lot"]["sizePcs"] == 6820 and r.data["summary"]["ncMajor"] == 1
    assert r.data["details"]["maxPower"]["stats"]["maximum"]["diffPct"] == 0.95  # Maximum row comes first in this report
    warn = [i.message for i in r.issues if i.code == "SAMPLING_BELOW_PLAN"]
    assert len(warn) == 2 and any("Static" in m for m in warn) and any("Dynamic" in m for m in warn)


def _walk(o):
    yield o
    if isinstance(o, dict):
        for k, v in o.items():
            yield k
            yield from _walk(v)
    elif isinstance(o, (list, tuple)):
        for v in o:
            yield from _walk(v)


def test_no_media_is_read_or_returned():
    for name in (L1, L2):
        items = list(_walk(parse(name).data))
        assert not [x for x in items if isinstance(x, (bytes, bytearray))]
        assert not [x for x in items if isinstance(x, str) and x.lower() in ("photos", "photo", "media", "images")]
    assert "Media in the document was ignored" in parse(L1).issues[0].message


def test_wrong_type_and_garbage():
    assert "WRONG_REPORT_TYPE" in codes(parse_pre_shipment((FIX / "ps_2026-09-30.docx").read_bytes()), "ERROR")  # a production report
    for payload in (b"", b"nope", b"PK\x03\x04junk"):
        r = parse_pre_shipment(payload)
        assert not r.ok and r.errors
    buf = io.BytesIO()
    Document().save(buf)
    assert "WRONG_REPORT_TYPE" in codes(parse_pre_shipment(buf.getvalue()), "ERROR")


def test_missing_lot_number_and_missing_results_are_errors():
    d = Document(FIX / L1)
    t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Report Number")
    t.rows[3].cells[1].paragraphs[0].runs[-1].text = "xxxxx"
    out = io.BytesIO()
    d.save(out)
    assert "LOT_NUMBER_MISSING" in codes(parse_pre_shipment(out.getvalue()), "ERROR")
    d = Document(FIX / L1)
    for t in list(d.tables):
        if "Testing Standard" in " ".join(c.text for c in t.rows[0].cells):
            t._tbl.getparent().remove(t._tbl)
    out = io.BytesIO()
    d.save(out)
    assert "RESULT_SUMMARY_MISSING" in codes(parse_pre_shipment(out.getvalue()), "ERROR")


# --------------------------------------------------------------------------- end to end (MongoDB)
from test_api import _mongo_up, client, login, upload_and_wait  # noqa: E402,F401

needs_mongo = pytest.mark.skipif(not _mongo_up(), reason="MongoDB not reachable")


@needs_mongo
async def test_upload_dashboard_reports_export(client):
    c, M = client, "pre_shipment"
    admin = await login(c, "admin@test.io", "AdminPass123")
    root = Path(os.environ["STORAGE_LOCAL_DIR"])
    before = {p for p in root.rglob("*") if p.is_file()} if root.exists() else set()

    for name in (L1, L2):
        up = await upload_and_wait(c, admin, name, (FIX / name).read_bytes(), module=M)
        assert up["status"] in ("SUCCESS", "SUCCESS_WITH_WARNINGS"), up["logs"]
    assert (await upload_and_wait(c, admin, L1, (FIX / L1).read_bytes(), module=M))["status"] == "FAILED"  # duplicate
    assert (await upload_and_wait(c, admin, L1, (FIX / L1).read_bytes(), module=M, replace="true"))["status"] != "FAILED"
    wrong = await upload_and_wait(c, admin, "x.docx", (FIX / "ps_2026-09-30.docx").read_bytes(), module=M)
    assert wrong["status"] == "FAILED" and any(x["code"] == "WRONG_REPORT_TYPE" for x in wrong["logs"])
    assert (await upload_and_wait(c, admin, "y.docx", (FIX / L1).read_bytes(), module="production_supervision"))["status"] == "FAILED"

    w = (await c.get(f"/{M}/dashboard", headers=admin)).json()["widgets"]
    assert w["kpi_lots"] == {"value": 2, "modules": 15066, "mwp": 9.2653}
    assert w["kpi_pass_rate"] == {"value": 100.0, "passed": 24, "tests": 24}
    assert w["kpi_ncs"] == {"value": 2, "major": 1, "minor": 1}
    assert w["kpi_power_dev"]["value"] == -0.122 and w["kpi_power_dev"]["outOfTolerance"] == 0
    assert [d["date"] for d in w["lot_sizes"]["data"]] == ["Lot 01", "Lot 02"] and w["lot_sizes"]["data"][0]["Modules"] == 8246
    assert [r["lot"] for r in w["lots"]["rows"]] == ["01", "02"] and w["lots"]["rows"][0]["tests"] == "12 / 12 passed"
    assert len(w["results"]["rows"]) == 24 and len(w["nonconformities"]["rows"]) == 2
    assert {r["lot"] for r in w["max_power"]["rows"]} == {"01", "02"} and w["max_power"]["rows"][0]["modules"] == 80
    assert len(w["equipment"]["rows"]) == 18 and all(r["status"] == "OK" for r in w["equipment"]["rows"])

    f = (await c.get(f"/{M}/filters", headers=admin)).json()
    assert f["lot"] == ["01", "02"] and len(f["project"]) == 1
    one = (await c.get(f"/{M}/dashboard?lot=02", headers=admin)).json()["widgets"]
    assert one["kpi_lots"]["modules"] == 6820 and len(one["results"]["rows"]) == 12
    assert (await c.get(f"/{M}/dashboard?date_from=2026-10-06&date_to=2026-10-31", headers=admin)).json()["widgets"]["kpi_lots"]["value"] == 1
    assert (await c.get(f"/{M}/dashboard?date_from=2026-10-09&date_to=2026-10-01", headers=admin)).status_code == 422

    assert (await c.put(f"/admin/widgets/{M}/kpi_ncs", headers=admin, json={"visible": False})).status_code == 200
    assert (await c.put(f"/admin/widgets/{M}/results", headers=admin, json={"columns": {"standard": False}})).status_code == 200
    w2 = (await c.get(f"/{M}/dashboard", headers=admin)).json()["widgets"]
    assert "kpi_ncs" not in w2 and "standard" not in w2["results"]["rows"][0] and "test" in w2["results"]["rows"][0]

    rep = (await c.get(f"/{M}/reports", headers=admin)).json()
    assert rep["total"] == 2 and rep["items"][0]["lotNo"] == "02"
    assert (await c.get(f"/{M}/reports?lot=01", headers=admin)).json()["total"] == 1
    assert (await c.get(f"/{M}/reports?q=2026-10-07", headers=admin)).json()["items"][0]["lotNo"] == "02"
    detail = (await c.get(f"/{M}/reports/{rep['items'][0]['id']}", headers=admin)).json()
    assert detail["details"]["maxPower"]["summary"]["modules"] == 80 and len(detail["results"]) == 12
    x = await c.get(f"/{M}/reports/export?format=xlsx", headers=admin)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content), read_only=True)
    assert {"Reports", "Results", "Non-conformities", "Equipment", "Max power", "Low irradiance"} <= set(wb.sheetnames)
    assert len(list(wb["Max power"].iter_rows(values_only=True))) == 1 + 160
    csv_ = await c.get(f"/{M}/reports/export?format=csv&lot=01", headers=admin)
    assert csv_.text.count("DSR-PSI-RES-2026-10-04") == 12
    assert (await c.get(f"/{M}/reports/export?q=nomatch", headers=admin)).status_code == 404

    assert ({p for p in root.rglob("*") if p.is_file()} if root.exists() else set()) == before  # nothing (no media, no .docx) stored

    ups = (await c.get("/uploads", headers=admin)).json()
    target = next(u for u in ups if u["module"] == M and u["status"].startswith("SUCCESS") and u["summary"]["reportNo"] == "DSR-PSI-RES-2026-10-07")
    assert (await c.delete(f"/uploads/{target['id']}", headers=admin)).status_code == 204
    assert (await c.get(f"/{M}/reports", headers=admin)).json()["total"] == 1
