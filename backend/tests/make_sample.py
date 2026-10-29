"""Builds tests/fixtures/sample_valid.docx by filling in the real DSR template."""
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).parent


def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def build(out: Path):
    d = Document(HERE / "fixtures" / "template.docx")
    for t in d.tables:
        first = t.rows[0].cells[0].text.strip()
        if first == "Report No.":
            for r, v in zip(t.rows, [": DSR-CLM-0001", ": 06-02-2026", ": Jinko Solar", ": Ningbo, China", ": Workshop 3"]):
                set_cell(r.cells[1], v)
        elif first == "INSPECTION CONCLUSION":
            set_cell(t.rows[1].cells[0], "Pass")
        elif first == "Date of Loading" and len(t.rows) == 4:
            for r in t.rows[2:]:
                for c in r.cells[:3]:
                    pass
            r = t.rows[2]
            vals = ["06-02-2026", "JKM550M-72HL4-BDVP", "1", "12", "12", "0", "", "8,640", "0", "4,752"]
            for c, v in zip(r.cells, vals):
                if v:
                    set_cell(c, v)
            # remove second sample row
            t._tbl.remove(t.rows[3]._tr)
    d.save(out)


if __name__ == "__main__":
    build(HERE / "fixtures" / "sample_valid.docx")
