"""Production Supervision: parser tests (always run) + end-to-end API test (needs MongoDB, see test_api.py)."""
import asyncio
import io
import os
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.modules.production_supervision.parser import parse_production_supervision

FIX = Path(__file__).parent / "fixtures"
A, B, C = "ps_2026-09-30.docx", "ps_2026-10-01.docx", "ps_2026-10-02.docx"


def parse(name):
    return parse_production_supervision((FIX / name).read_bytes())


def codes(r, level):
    return {i.code for i in r.issues if i.level == level}


def test_report_30_sep():
    r = parse(A)
    assert r.ok, [i.message for i in r.errors]
    d, rep = r.data, r.data["report"]
    assert rep["reportNo"] == "DSR-PS-RES-2026-09-30" and rep["reportDate"].day == 30
    assert rep["spvName"].startswith("Solapur Ganeko One") and rep["projectName"].startswith("Ganeko Two")
    assert [(p["date"].day, p["powerBin"], p["actual"]["qty"], p["planned"]["qty"]) for p in d["production"]] == \
        [(28, 615, 1416, 2000), (29, 615, 4120, 5000), (30, 615, 3912, 4600)]
    assert d["productionTotal"]["actual"]["qty"] == 9448
    q = {x["kind"]: x for x in d["qualification"]}
    assert q["daily"]["qualified"]["615"]["qty"] == 3689 and q["cumulative"]["rejected"]["qty"] == 72 and q["cumulative"]["suspected"]["qty"] == 126
    assert len(d["bom"]) == 2 and {(b["line"], b["shift"]) for b in d["bom"]} == {("2", "Day"), ("2", "Night")}
    assert len(d["ipqc"]) == 56 and len(d["flasherResults"]) == 6 and len(d["nonConformities"]) == 3
    assert d["summary"]["actualQty"] == 9448 and d["summary"]["openNcs"] == 0
    assert d["gelOverall"] == "Pass" and d["dimensionOverall"] == "Pass"


def test_open_nc_and_bom_not_recorded_handling():
    b = parse(B).data
    assert b["summary"]["actualQty"] == 13758 and b["summary"]["openNcs"] == 1  # 'Open' wins over the record's odd 'Yes'
    assert [(bb["line"], bb["shift"]) for bb in b["bom"]] == [("1 & 2", "Day"), ("1 & 2", "Night")]
    c = parse(C)
    assert c.ok and c.data["summary"]["actualQty"] == 17328 and c.data["summary"]["days"] == 5
    night = next(x for x in c.data["bom"] if x["shift"] == "Night")
    assert all(m["makeMatches"] is None for m in night["materials"])  # '-' = not recorded, not a deviation
    assert c.data["summary"]["bomDiffer"] == 2


def _walk(o):
    yield o
    if isinstance(o, dict):
        for k, v in o.items():
            yield k
            yield from _walk(v)
    elif isinstance(o, (list, tuple)):
        for v in o:
            yield from _walk(v)


def test_no_media_is_read_or_returned():
    for name in (A, B, C):
        items = list(_walk(parse(name).data))
        assert not [x for x in items if isinstance(x, (bytes, bytearray))]  # no image/video blobs
        assert not [x for x in items if isinstance(x, str) and x.lower() in ("photos", "photo", "media", "images")]
    assert "Media in the document was ignored" in parse(A).issues[0].message


def test_wrong_document_type_is_rejected_with_clear_error():
    r = parse_production_supervision((FIX / "premier_report_slim.docx").read_bytes())  # a container loading report
    assert not r.ok and "WRONG_REPORT_TYPE" in codes(r, "ERROR")


@pytest.mark.parametrize("payload", [b"", b"nope", b"PK\x03\x04junk"])
def test_garbage_never_raises(payload):
    r = parse_production_supervision(payload)
    assert not r.ok and r.errors


def test_blank_docx_and_missing_snapshot():
    buf = io.BytesIO()
    Document().save(buf)
    assert "WRONG_REPORT_TYPE" in codes(parse_production_supervision(buf.getvalue()), "ERROR")
    d = Document(FIX / A)
    for t in list(d.tables):  # remove the production snapshot table
        if "Planned Output" in " ".join(c.text for c in t.rows[0].cells):
            t._tbl.getparent().remove(t._tbl)
    out = io.BytesIO()
    d.save(out)
    assert "PRODUCTION_TABLE_MISSING" in codes(parse_production_supervision(out.getvalue()), "ERROR")


def test_placeholder_report_no_is_an_error():
    d = Document(FIX / A)
    t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Report No.")
    t.rows[0].cells[1].paragraphs[0].runs[-1].text = "xxxxx"
    out = io.BytesIO()
    d.save(out)
    assert "REPORT_NO_MISSING" in codes(parse_production_supervision(out.getvalue()), "ERROR")


# --------------------------------------------------------------------------- end to end (MongoDB)
from test_api import _mongo_up, client, login, upload_and_wait  # noqa: E402,F401  (reuse fixtures/helpers)

needs_mongo = pytest.mark.skipif(not _mongo_up(), reason="MongoDB not reachable")


@needs_mongo
async def test_upload_dashboard_reports_and_export(client):
    c = client
    admin = await login(c, "admin@test.io", "AdminPass123")
    M = "production_supervision"
    root = Path(os.environ["STORAGE_LOCAL_DIR"])
    stored_before = {p for p in root.rglob("*") if p.is_file()} if root.exists() else set()

    # uploading the three cumulative reports in order; later reports repeat earlier days
    for name in (A, B, C):
        up = await upload_and_wait(c, admin, name, (FIX / name).read_bytes(), module=M)
        assert up["status"] in ("SUCCESS", "SUCCESS_WITH_WARNINGS"), up["logs"]
    assert (await upload_and_wait(c, admin, A, (FIX / A).read_bytes(), module=M))["status"] == "FAILED"  # duplicate
    assert (await upload_and_wait(c, admin, A, (FIX / A).read_bytes(), module=M, replace="true"))["status"] != "FAILED"
    # a container report sent to the wrong module fails cleanly, as does the reverse
    wrong = await upload_and_wait(c, admin, "x.docx", (FIX / "premier_report_slim.docx").read_bytes(), module=M)
    assert wrong["status"] == "FAILED" and any(x["code"] == "WRONG_REPORT_TYPE" for x in wrong["logs"])
    assert (await upload_and_wait(c, admin, "p.docx", (FIX / C).read_bytes(), module="container_loading"))["status"] == "FAILED"

    w = (await c.get(f"/{M}/dashboard", headers=admin)).json()["widgets"]
    # latest report (02-Oct) wins for the overlapping days; 5 production days in total
    assert w["kpi_output"]["value"] == 17328 and w["kpi_achievement"]["planned"] == 20200 and w["kpi_achievement"]["value"] == 85.8
    assert w["kpi_rejected"] == {"value": 437, "rejected": 147, "suspected": 290}
    assert w["kpi_open_nc"] == {"value": 1, "total": 4}
    assert [d["date"] for d in w["production_trend"]["data"]] == ["2026-09-28", "2026-09-29", "2026-09-30", "2026-10-01", "2026-10-02"]
    assert len(w["daily_production"]["rows"]) == 5 and w["daily_production"]["rows"][0]["date"] == "2026-10-02"
    q = w["qualification"]["rows"]
    assert q[-1]["date"] == "Cumulative" and q[-1]["qualified"]["615"]["qty"] == 15314
    assert any(r["status"] == "OPEN" for r in w["nonconformities"]["rows"])
    assert w["quality_checks"]["source"] == "DSR-PS-RES-2026-10-02" and len(w["quality_checks"]["rows"]) == 10

    # filters
    f = (await c.get(f"/{M}/filters", headers=admin)).json()
    assert f["power_bin"] == ["615"] and len(f["project"]) == 1
    one = (await c.get(f"/{M}/dashboard?date_from=2026-09-30&date_to=2026-09-30", headers=admin)).json()["widgets"]
    assert one["kpi_output"]["value"] == 3912
    assert (await c.get(f"/{M}/dashboard?power_bin=610", headers=admin)).json()["widgets"]["kpi_output"]["value"] == 0
    assert (await c.get(f"/{M}/dashboard?date_from=2026-10-05&date_to=2026-10-01", headers=admin)).status_code == 422

    # widget/column admin applies to this module too
    assert (await c.put(f"/admin/widgets/{M}/kpi_open_nc", headers=admin, json={"visible": False})).status_code == 200
    assert (await c.put(f"/admin/widgets/{M}/daily_production", headers=admin, json={"columns": {"night_shift": False}})).status_code == 200
    w2 = (await c.get(f"/{M}/dashboard", headers=admin)).json()["widgets"]
    assert "kpi_open_nc" not in w2 and "night_shift" not in w2["daily_production"]["rows"][0]

    # reports tab
    rep = (await c.get(f"/{M}/reports", headers=admin)).json()
    assert rep["total"] == 3 and rep["items"][0]["reportNo"] == "DSR-PS-RES-2026-10-02"
    assert (await c.get(f"/{M}/reports?q=2026-10-01", headers=admin)).json()["total"] == 1
    assert (await c.get(f"/{M}/reports?date_from=2026-10-02", headers=admin)).json()["total"] == 1  # only the report covering 02-Oct
    detail = (await c.get(f"/{M}/reports/{rep['items'][0]['id']}", headers=admin)).json()
    assert detail["details"]["quality"] and len(detail["production"]) == 5
    x = await c.get(f"/{M}/reports/export?format=xlsx", headers=admin)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content), read_only=True)
    assert {"Reports", "Production", "Qualification", "Non-conformities", "BOM", "IPQC"} <= set(wb.sheetnames)
    assert len(list(wb["Reports"].iter_rows(values_only=True))) == 1 + 3
    csv_ = await c.get(f"/{M}/reports/export?format=csv&q=2026-09-30", headers=admin)
    assert csv_.text.count("DSR-PS-RES-2026-09-30") == 3
    assert (await c.get(f"/{M}/reports/export?q=nomatch", headers=admin)).status_code == 404

    # media / originals are never stored
    stored_after = {p for p in root.rglob("*") if p.is_file()} if root.exists() else set()
    assert stored_after == stored_before

    # deleting an upload removes everything it wrote
    ups = (await c.get("/uploads", headers=admin)).json()
    target = next(u for u in ups if u["module"] == M and u["status"].startswith("SUCCESS") and u["summary"]["reportNo"] == "DSR-PS-RES-2026-10-02")
    assert (await c.delete(f"/uploads/{target['id']}", headers=admin)).status_code == 204
    assert (await c.get(f"/{M}/reports", headers=admin)).json()["total"] == 2
    assert (await c.get(f"/{M}/dashboard", headers=admin)).json()["widgets"]["kpi_output"]["value"] == 13758
