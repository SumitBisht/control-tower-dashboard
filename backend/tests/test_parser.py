import io
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.modules.container_loading.parser import parse_container_loading, parse_date, status_of

FIX = Path(__file__).parent / "fixtures"


def load(name):
    return (FIX / name).read_bytes()


def codes(r, level=None):
    return {i.code for i in r.issues if level is None or i.level == level}


@pytest.fixture(scope="session", autouse=True)
def sample():
    from make_sample import build  # noqa: PLC0415
    build(FIX / "sample_valid.docx")


def test_unfilled_template_is_rejected_with_clear_errors():
    r = parse_container_loading(load("template.docx"))
    assert not r.ok
    assert {"REPORT_NO_MISSING", "REPORT_DATE_INVALID"} <= codes(r, "ERROR")


def test_valid_report_parses():
    r = parse_container_loading(load("sample_valid.docx"))
    assert r.ok, [i.message for i in r.errors]
    cs = r.data["containers"]
    assert len(cs) == 12 and sum(c["moduleQty"] for c in cs) == 8640
    assert cs[0]["psiLotNo"] == "1" and cs[0]["capacityKwp"] == 396.0
    assert cs[0]["inspection"]["checklist"]["shockLabel"] == "NA"
    assert r.data["report"]["conclusion"] == "PASS"
    assert "SEAL_MISMATCH" in codes(r, "WARNING")  # the template itself contains mismatching seals


@pytest.mark.parametrize("payload", [b"", b"not a zip", b"PK\x03\x04garbage"])
def test_garbage_never_raises(payload):
    r = parse_container_loading(payload)
    assert not r.ok and r.errors


def test_zip_without_word_document_is_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    r = parse_container_loading(buf.getvalue())
    assert not r.ok


def test_blank_document_reports_missing_structure():
    buf = io.BytesIO()
    Document().save(buf)
    r = parse_container_loading(buf.getvalue())
    assert "REPORT_INFO_MISSING" in codes(r, "ERROR")


def test_bad_container_row_collects_all_errors():
    d = Document(FIX / "sample_valid.docx")
    t = next(t for t in d.tables if len(t.rows) > 3 and "Containers No." in " ".join(c.text for c in t.rows[0].cells))
    t.rows[2].cells[7].paragraphs[0].runs[0].text = "abc"      # bad qty
    t.rows[3].cells[9].paragraphs[0].runs[0].text = "31-31-2026"  # bad date
    buf = io.BytesIO()
    d.save(buf)
    r = parse_container_loading(buf.getvalue())
    assert not r.ok and {"MODULE_QTY_INVALID", "LOADING_DATE_INVALID"} <= codes(r, "ERROR")


def test_helpers():
    assert parse_date("06.02.2026").month == 2 and parse_date("5-Sep-2026").month == 9 and parse_date("xxxxx") is None
    assert status_of("No", "shockLabel") == "NA" and status_of("OK", "weather") == "OK" and status_of("broken", "weather") == "?"


def test_premier_report_real_layout():
    """Real DSR report: Client/Project/SPV table, no seal column, 'Container Number N' headings, PSI lots like '1, 2'."""
    r = parse_container_loading(load("premier_report_slim.docx"))
    assert r.ok, [i.message for i in r.errors]
    rep, cs = r.data["report"], r.data["containers"]
    assert rep["reportNo"] == "DSR-CLC-RES-2026-10-11" and rep["reportDate"].day == 11
    assert rep["spvName"].startswith("Solapur Ganeko One") and rep["projectName"].startswith("Ganeko Two")
    assert len(cs) == 25 and sum(c["moduleQty"] for c in cs) == 15066
    assert round(sum(c["capacityKwp"] for c in cs), 2) == 9265.59
    assert {c["psiLotNo"] for c in cs} == {"1", "1, 2", "2"} and ["1", "2"] in [c["psiLots"] for c in cs]
    assert all(c["inspection"]["result"] == "PASS" for c in cs)
    assert cs[0]["shipment"]["truckDeparture"].hour == 17
    assert sum(len(c["photos"]) for c in cs) == 196
    assert not codes(r, "WARNING") - {"SUMMARY_MISMATCH"}  # clean report: no warnings
